import io
import json
import zipfile

import pymupdf
import pytest
from docx import Document
from fastapi.testclient import TestClient

from app.lexicon import summarize
from app.lexicon.extract import ExtractError, extract
from app.main import app

client = TestClient(app)

PARA = "The quick analysis of systems reveals many useful ideas about design. " * 30  # ~330 words


def make_txt() -> bytes:
    toc = "Contents\nChapter 1 The Start\nChapter 2 The Middle\nChapter 3 The End\n\n"
    body = ""
    for n, name in [(1, "The Start"), (2, "The Middle"), (3, "The End")]:
        body += f"\nChapter {n} {name}\n{PARA}\n"
        # a running page header in the middle of the chapter, and a cross-reference
        body += f"\nChapter {n} {name}\n{PARA}\nChapter 3 explains the rest.\n"
    return (toc + body).encode()


def make_docx() -> bytes:
    doc = Document()
    for name in ["Alpha", "Beta"]:
        doc.add_heading(name, level=1)
        doc.add_paragraph(PARA)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def make_epub() -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("mimetype", "application/epub+zip")
        z.writestr(
            "META-INF/container.xml",
            '<container><rootfiles><rootfile full-path="OEBPS/content.opf"/></rootfiles></container>',
        )
        z.writestr(
            "OEBPS/content.opf",
            '<package><metadata><dc:title xmlns:dc="http://purl.org/dc/elements/1.1/">Test Book</dc:title></metadata>'
            '<manifest><item id="cover" href="cover.xhtml"/><item id="c1" href="c1.xhtml"/><item id="c2" href="c2.xhtml"/></manifest>'
            '<spine><itemref idref="cover"/><itemref idref="c1"/><itemref idref="c2"/></spine></package>',
        )
        z.writestr("OEBPS/cover.xhtml", "<html><body><h1>Cover</h1></body></html>")
        z.writestr("OEBPS/c1.xhtml", f"<html><body><h1>One</h1><p>{PARA}</p></body></html>")
        z.writestr("OEBPS/c2.xhtml", f"<html><body><h1>Two</h1><p>{PARA}</p></body></html>")
    return buf.getvalue()


def make_pdf_with_big_headings() -> bytes:
    doc = pymupdf.open()
    for name in ["Origins", "Growth", "Decline"]:
        page = doc.new_page()
        page.insert_text((72, 80), name, fontsize=28)
        page.insert_textbox(pymupdf.Rect(72, 120, 540, 780), PARA, fontsize=10)
        page = doc.new_page()
        page.insert_textbox(pymupdf.Rect(72, 72, 540, 780), PARA, fontsize=10)
    return doc.tobytes()


def test_txt_skips_toc_running_headers_and_references():
    book = extract("novel.txt", make_txt())
    assert [c.title for c in book.chapters] == ["Chapter 1 The Start", "Chapter 2 The Middle", "Chapter 3 The End"]


def test_docx_uses_heading_styles():
    assert [c.title for c in extract("notes.docx", make_docx()).chapters] == ["Alpha", "Beta"]


def test_epub_follows_spine_and_drops_short_pages():
    book = extract("book.epub", make_epub())
    assert book.title == "Test Book"
    assert [c.title for c in book.chapters] == ["One", "Two"]


def test_pdf_without_bookmarks_uses_font_size():
    book = extract("book.pdf", make_pdf_with_big_headings())
    assert [c.title for c in book.chapters] == ["Origins", "Growth", "Decline"]


def test_unstructured_text_is_chunked():
    book = extract("blob.txt", (PARA * 20).encode())
    assert len(book.chapters) >= 2
    assert book.chapters[0].title == "Part 1"


@pytest.mark.parametrize(
    ("name", "data", "message"),
    [
        ("book.mobi", b"x", "unsupported format"),
        ("book.pdf", b"not a pdf", "could not read"),
        ("tiny.txt", b"hello", "no readable text"),
    ],
)
def test_bad_files_give_readable_errors(name, data, message):
    with pytest.raises(ExtractError, match=message):
        extract(name, data)


def test_upload_then_stream_summary(monkeypatch):
    async def fake_summarize(book, chapter, text, on_model=None):
        yield "## Summary\n"
        yield f"about {chapter}"
        if on_model:
            on_model("fake-model")

    monkeypatch.setattr(summarize, "summarize", fake_summarize)
    monkeypatch.setattr(summarize, "ready", lambda: True)

    res = client.post("/api/lexicon/books", files={"file": ("notes.docx", make_docx())})
    assert res.status_code == 200
    book = res.json()
    assert [c["title"] for c in book["chapters"]] == ["Alpha", "Beta"]

    res = client.post(f"/api/lexicon/books/{book['id']}/chapters/1/summary")
    assert res.status_code == 200
    text, meta = res.text.split("\n[[meta]]")
    assert text == "## Summary\nabout Beta"
    assert json.loads(meta)["model"] == "fake-model"


def test_upload_rejects_bad_file():
    res = client.post("/api/lexicon/books", files={"file": ("x.exe", b"MZ")})
    assert res.status_code == 422
    assert "unsupported" in res.json()["detail"]


def test_summary_without_key_is_503(monkeypatch):
    monkeypatch.setattr(summarize, "ready", lambda: False)
    res = client.post("/api/lexicon/books/whatever/chapters/0/summary")
    assert res.status_code == 503


def test_falls_back_to_next_model_when_one_is_busy(monkeypatch):
    import asyncio

    from google.genai import errors

    from app.core.config import settings

    calls = []

    class FakeModels:
        async def generate_content_stream(self, model, contents, config):
            calls.append(model)
            if model == "busy-model":
                raise errors.ServerError(503, {"error": {"code": 503, "status": "UNAVAILABLE", "message": "busy"}})

            async def gen():
                class Chunk:
                    text = "answer"

                yield Chunk()

            return gen()

    class FakeClient:
        class aio:
            models = FakeModels()

    monkeypatch.setattr(settings, "gemini_models", ["busy-model", "good-model"])
    monkeypatch.setattr(summarize, "_get_client", lambda: FakeClient())

    async def collect():
        return [piece async for piece in summarize.summarize("b", "c", "text")]

    assert asyncio.run(collect()) == ["answer"]
    assert calls == ["busy-model", "good-model"]


def test_restarts_on_next_model_when_one_dies_mid_answer(monkeypatch):
    import asyncio

    from google.genai import errors

    from app.core.config import settings

    class FakeModels:
        async def generate_content_stream(self, model, contents, config):
            async def gen():
                class Chunk:
                    text = "half an answer" if model == "flaky-model" else "whole answer"

                yield Chunk()
                if model == "flaky-model":
                    raise errors.ServerError(503, {"error": {"code": 503, "status": "UNAVAILABLE", "message": "busy"}})

            return gen()

    class FakeClient:
        class aio:
            models = FakeModels()

    monkeypatch.setattr(settings, "gemini_models", ["flaky-model", "good-model"])
    monkeypatch.setattr(summarize, "_get_client", lambda: FakeClient())

    async def collect():
        return [piece async for piece in summarize.summarize("b", "c", "text")]

    # The page keeps only what follows the last marker.
    assert asyncio.run(collect()) == ["half an answer", summarize.RESTART_MARK, "whole answer"]


def test_a_model_that_goes_quiet_is_abandoned(monkeypatch):
    import asyncio

    from app.core.config import settings

    class FakeModels:
        async def generate_content_stream(self, model, contents, config):
            async def gen():
                class Chunk:
                    text = "half an answer" if model == "stuck-model" else "whole answer"

                yield Chunk()
                if model == "stuck-model":
                    await asyncio.sleep(60)  # never finishes

            return gen()

    class FakeClient:
        class aio:
            models = FakeModels()

    monkeypatch.setattr(settings, "gemini_models", ["stuck-model", "good-model"])
    monkeypatch.setattr(summarize, "_get_client", lambda: FakeClient())
    monkeypatch.setattr(summarize, "STALL_SECONDS", 0.05)

    async def collect():
        return [piece async for piece in summarize.summarize("b", "c", "text")]

    assert asyncio.run(collect()) == ["half an answer", summarize.RESTART_MARK, "whole answer"]
